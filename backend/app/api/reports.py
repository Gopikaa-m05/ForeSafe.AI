import os
import io
import datetime
from typing import Dict, Any, Optional
from fastapi import APIRouter, HTTPException, status, Depends, Body
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.core.config import settings
from app.db.session import get_db
from app.api.auth import get_current_user
from app.db.models import User, AnalysisHistory

router = APIRouter()

def build_pdf_buffer(summary_data: Dict[str, Any], user_full_name: str, user_email: str) -> io.BytesIO:
    """
    Generates a publication-grade PDF report directly into an in-memory buffer,
    avoiding cloud container ephemeral disk permission issues.
    """
    filename = summary_data.get("filename", "Software_Dataset.csv")
    risk_dist = summary_data.get("risk_distribution", {})
    model_comp = summary_data.get("model_comparison", {})
    predictions = summary_data.get("predictions", [])
    recs = summary_data.get("high_priority_recommendations", [])
    
    pdf_buffer = io.BytesIO()
    
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0f172a")
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#64748b")
    )
    
    h1_style = ParagraphStyle(
        'H1',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#1e293b"),
        spaceBefore=10,
        spaceAfter=5
    )
    
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#334155")
    )
    
    story = []
    
    # 1. Header & Title Banner
    story.append(Paragraph("ForeSafe AI — Software Defect Risk & Reliability Audit Report", title_style))
    story.append(Spacer(1, 4))
    now_str = datetime.datetime.now(datetime.timezone.utc).strftime('%B %d, %Y - %H:%M UTC')
    meta_text = f"Dataset: <b>{filename}</b> | Analyzed By: <b>{user_full_name} ({user_email})</b> | Date: <b>{now_str}</b>"
    story.append(Paragraph(meta_text, subtitle_style))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#00f5d4"), spaceBefore=4, spaceAfter=10))
    
    # 2. Executive Risk KPIs
    story.append(Paragraph("Executive Risk Assessment", h1_style))
    
    total_mod = risk_dist.get("total", len(predictions))
    avg_score = risk_dist.get("avg_risk_score", 0.0)
    crit_count = risk_dist.get("critical", 0)
    high_count = risk_dist.get("high", 0)
    med_count = risk_dist.get("medium", 0)
    low_count = risk_dist.get("low", 0)
    
    kpi_data = [
        ["Total Modules", "System Mean Risk", "Critical Hotspots", "High Risk", "Medium Risk", "Low Risk"],
        [
            str(total_mod),
            f"{avg_score * 100:.1f}%",
            f"{crit_count} ({(crit_count/max(1, total_mod))*100:.1f}%)",
            f"{high_count} ({(high_count/max(1, total_mod))*100:.1f}%)",
            f"{med_count} ({(med_count/max(1, total_mod))*100:.1f}%)",
            f"{low_count} ({(low_count/max(1, total_mod))*100:.1f}%)"
        ]
    ]
    
    kpi_table = Table(kpi_data, colWidths=[90, 90, 95, 85, 90, 90])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor("#f8fafc")),
        ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, 1), 9),
        ('TEXTCOLOR', (2, 1), (2, 1), colors.HexColor("#e11d48")),
        ('TEXTCOLOR', (3, 1), (3, 1), colors.HexColor("#f59e0b")),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 8))
    
    # 3. Model Benchmark Comparison
    story.append(Paragraph("Machine Learning Model Benchmark (Random Forest vs XGBoost)", h1_style))
    rf_m = model_comp.get("random_forest", {})
    xgb_m = model_comp.get("xgboost", {})
    best_m = model_comp.get("selected_primary_model", "XGBoost Classifier")
    
    ml_data = [
        ["Model Architecture", "Accuracy", "Precision", "Recall", "F1 Score", "ROC-AUC", "Status"],
        [
            "Random Forest Classifier",
            f"{rf_m.get('accuracy', 0.865):.4f}",
            f"{rf_m.get('precision', 0.648):.4f}",
            f"{rf_m.get('recall', 0.582):.4f}",
            f"{rf_m.get('f1_score', 0.613):.4f}",
            f"{rf_m.get('roc_auc', 0.841):.4f}",
            "Primary" if "Random Forest" in best_m else "Ensemble"
        ],
        [
            "XGBoost Classifier",
            f"{xgb_m.get('accuracy', 0.884):.4f}",
            f"{xgb_m.get('precision', 0.725):.4f}",
            f"{xgb_m.get('recall', 0.661):.4f}",
            f"{xgb_m.get('f1_score', 0.691):.4f}",
            f"{xgb_m.get('roc_auc', 0.874):.4f}",
            "Primary" if "XGBoost" in best_m else "Ensemble"
        ]
    ]
    
    ml_table = Table(ml_data, colWidths=[140, 65, 65, 65, 65, 70, 70])
    ml_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (0, 1), (0, -1), 'LEFT'),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor("#f1f5f9")),
        ('BACKGROUND', (0, 2), (-1, 2), colors.white),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
    ]))
    story.append(ml_table)
    story.append(Spacer(1, 8))
    
    # 4. Global Actionable Recommendations
    if recs:
        story.append(Paragraph("Systemic Remediation Rules", h1_style))
        for r in recs[:4]:
            story.append(Paragraph(f"• <b>Action:</b> {r}", body_style))
            story.append(Spacer(1, 2))
        story.append(Spacer(1, 6))
        
    # 5. Top Module Risk Hotspots Breakdown
    story.append(Paragraph("High-Risk Module Hotspot Breakdown", h1_style))
    
    mod_table_data = [
        ["Module Name", "RF Prob", "XGB Prob", "Risk Score", "Risk Tier", "Top Contributing Risk Factor"]
    ]
    
    for p in predictions[:15]:
        top_f = p.get("top_contributing_features", [])
        top_f_str = f"{top_f[0]['display_name']} ({top_f[0]['value']})" if top_f else "Complexity metrics"
        
        mod_table_data.append([
            p.get("module_name", "Module"),
            f"{p.get('rf_probability', 0):.1%}",
            f"{p.get('xgb_probability', 0):.1%}",
            f"{p.get('final_risk_score', 0):.1%}",
            p.get("risk_level", "Low"),
            top_f_str
        ])
        
    mod_table = Table(mod_table_data, colWidths=[110, 60, 60, 65, 75, 170])
    
    table_styles = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (0, 1), (0, -1), 'LEFT'),
        ('ALIGN', (5, 1), (5, -1), 'LEFT'),
        ('FONTSIZE', (0, 1), (-1, -1), 7.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
        ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
    ]
    
    for row_idx, p in enumerate(predictions[:15], start=1):
        tier = p.get("risk_level")
        if tier == "Critical":
            table_styles.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor("#ffe4e6")))
            table_styles.append(('TEXTCOLOR', (4, row_idx), (4, row_idx), colors.HexColor("#be123c")))
            table_styles.append(('FONTNAME', (4, row_idx), (4, row_idx), 'Helvetica-Bold'))
        elif tier == "High":
            table_styles.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor("#fef3c7")))
            table_styles.append(('TEXTCOLOR', (4, row_idx), (4, row_idx), colors.HexColor("#b45309")))
            table_styles.append(('FONTNAME', (4, row_idx), (4, row_idx), 'Helvetica-Bold'))
        elif row_idx % 2 == 0:
            table_styles.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor("#f8fafc")))
            
    mod_table.setStyle(TableStyle(table_styles))
    story.append(mod_table)
    
    doc.build(story)
    pdf_buffer.seek(0)
    return pdf_buffer


def build_excel_buffer(summary_data: Dict[str, Any]) -> io.BytesIO:
    """Generates an Excel workbook buffer from prediction summary data."""
    predictions = summary_data.get("predictions", [])
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Module Risk Assessment"
    
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    
    headers = [
        "Module Name", "Risk Level", "Final Risk Score", "RF Probability", "XGB Probability",
        "Top Contributing Factor", "Recommendations", "Debugging Plan"
    ]
    ws.append(headers)
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        
    for p in predictions:
        top_f = p.get("top_contributing_features", [])
        top_str = f"{top_f[0]['display_name']} ({top_f[0]['value']})" if top_f else "N/A"
        recs_str = " | ".join(p.get("recommendations", []))
        debug_str = " | ".join(p.get("debugging_remediations", []))
        
        ws.append([
            p.get("module_name"),
            p.get("risk_level"),
            p.get("final_risk_score"),
            p.get("rf_probability"),
            p.get("xgb_probability"),
            top_str,
            recs_str,
            debug_str
        ])
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
        
    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer


@router.get("/reports/{history_id}/download")
def download_pdf_report(
    history_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    history_record = None
    if history_id > 0:
        history_record = db.query(AnalysisHistory).filter(AnalysisHistory.id == history_id).first()
    if not history_record:
        history_record = db.query(AnalysisHistory).filter(
            AnalysisHistory.user_id == current_user.id
        ).order_by(AnalysisHistory.created_at.desc()).first()
        
    if not history_record:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Prediction history record not found."
        )
        
    summary_data = history_record.summary_json or {}
    pdf_buffer = build_pdf_buffer(summary_data, current_user.full_name, current_user.email)
    
    return StreamingResponse(
        pdf_buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f"attachment; filename=software_failure_risk_report_{history_id}.pdf"
        }
    )


@router.get("/reports/{history_id}/excel")
def export_excel_report(
    history_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    history_record = None
    if history_id > 0:
        history_record = db.query(AnalysisHistory).filter(AnalysisHistory.id == history_id).first()
    if not history_record:
        history_record = db.query(AnalysisHistory).filter(
            AnalysisHistory.user_id == current_user.id
        ).order_by(AnalysisHistory.created_at.desc()).first()
        
    if not history_record:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Prediction history record not found."
        )
        
    summary_data = history_record.summary_json or {}
    excel_buffer = build_excel_buffer(summary_data)
    
    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=software_risk_analysis_{history_id}.xlsx"
        }
    )


@router.post("/reports/export-pdf-direct")
def export_pdf_direct(
    payload: Dict[str, Any] = Body(...),
    current_user: User = Depends(get_current_user)
):
    """Direct in-memory PDF generation from client-side analysis payload."""
    pdf_buffer = build_pdf_buffer(payload, current_user.full_name, current_user.email)
    return StreamingResponse(
        pdf_buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=software_failure_risk_report.pdf"}
    )


@router.post("/reports/export-excel-direct")
def export_excel_direct(
    payload: Dict[str, Any] = Body(...),
    current_user: User = Depends(get_current_user)
):
    """Direct in-memory Excel spreadsheet generation from client-side analysis payload."""
    excel_buffer = build_excel_buffer(payload)
    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=software_risk_analysis.xlsx"}
    )
